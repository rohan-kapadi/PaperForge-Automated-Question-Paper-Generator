import os
import io
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="000000", sz="4"):
    """Set black borders on a cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name, border_style in [('top', top), ('left', left), ('bottom', bottom), ('right', right)]:
        if border_style:
            b = OxmlElement(f'w:{border_name}')
            b.set(qn('w:val'), border_style)
            b.set(qn('w:sz'), sz)
            b.set(qn('w:space'), '0')
            b.set(qn('w:color'), color)
            tcBorders.append(b)
        else:
            b = OxmlElement(f'w:{border_name}')
            b.set(qn('w:val'), 'none')
            tcBorders.append(b)
    tcPr.append(tcBorders)

def set_cell_margins(cell, top=50, bottom=50, left=80, right=80):
    """Set cell padding in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m_name, m_val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m_name}')
        node.set(qn('w:w'), str(m_val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def make_row_cant_split(row):
    """Ensure table row does not split across pages in Word."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:cantSplit'))

def set_cell_shading(cell, color="D9D9D9"):
    """Set background fill color on a cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)

def parse_marks_co_btl(val):
    """
    Format Marks/CO/BTL into 3 lines matching PCCOER standard:
    Line 1: (Marks-05)
    Line 2: [P251.1]
    Line 3: BT6
    """
    import re
    if not val:
        return [('(Marks-05)', False), ('[CO1]', True), ('BT2', True)]
    
    val_str = str(val).strip()
    lines = [l.strip() for l in val_str.split('\n') if l.strip()]
    if len(lines) == 3:
        return [(lines[0], False), (lines[1], True), (lines[2], True)]
    
    parts = [p.strip() for p in val_str.split('/')]
    m_raw = parts[0] if len(parts) > 0 else '5'
    digits = re.search(r'\d+', m_raw)
    if digits:
        marks_str = f'(Marks-{int(digits.group(0)):02d})'
    else:
        marks_str = f'({m_raw})'
    
    co_raw = parts[1] if len(parts) > 1 else 'CO1'
    co_clean = co_raw.strip('[] ')
    co_str = f'[{co_clean}]'
    
    bt_raw = parts[2] if len(parts) > 2 else 'BT2'
    bt_m = re.search(r'(?:BT|L)?\s*(\d+)', bt_raw, re.IGNORECASE)
    if bt_m:
        bt_str = f'BT{bt_m.group(1)}'
    else:
        bt_str = bt_raw
        
    return [(marks_str, False), (co_str, True), (bt_str, True)]

def generate_pccoer_docx_stream(paper_data: dict) -> io.BytesIO:
    """
    Generate the official PCCOER question paper format in DOCX.
    Returns a BytesIO stream.
    """
    doc = docx.Document()
    
    # Page Margins
    # A4 Page Size & Margins
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.32)
        section.bottom_margin = Inches(0.32)
        section.left_margin = Inches(0.4)
        section.right_margin = Inches(0.4)
        
        # Footer
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        f_p.paragraph_format.space_before = Pt(0)
        f_p.paragraph_format.space_after = Pt(0)
        rev_text = f"Rev.: {paper_data.get('rev', '00')}      Date: {paper_data.get('rev_date', '01-09-2025')}"
        r_f1 = f_p.add_run(rev_text)
        r_f1.font.name = 'Times New Roman'
        r_f1.font.size = Pt(10)
        
        r_f2 = f_p.add_run("\t\t\t\t\t\t\t\tPage 1 of 2")
        r_f2.font.name = 'Times New Roman'
        r_f2.font.size = Pt(10)

    base_dir = os.path.dirname(os.path.abspath(__file__))
    # Check possible logo paths
    logo_paths = [
        os.path.join(base_dir, '..', 'frontend', 'public', 'logos'),
        os.path.join(base_dir, 'logos'),
    ]
    pccoer_logo = None
    pcet_logo = None
    for lp in logo_paths:
        p1 = os.path.join(lp, 'pccoer_logo.jpg')
        p2 = os.path.join(lp, 'pcet_logo.jpg')
        if os.path.exists(p1) and os.path.exists(p2):
            pccoer_logo = p1
            pcet_logo = p2
            break

    # 1. Header Table (3 columns: Left Logo | Center Text | Right Logo)
    header_table = doc.add_table(rows=2, cols=3)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False

    col_widths = [Inches(1.4), Inches(4.67), Inches(1.4)]
    for row in header_table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    # Row 0: Logos & Trust Info
    c00 = header_table.cell(0, 0)
    c00.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = c00.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if pccoer_logo and os.path.exists(pccoer_logo):
        p.add_run().add_picture(pccoer_logo, width=Inches(0.95))

    c01 = header_table.cell(0, 1)
    c01.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = c01.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    
    r = p.add_run("Pimpri Chinchwad Education Trust's\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r.bold = True
    
    r = p.add_run("Pimpri Chinchwad College of Engineering & Research Ravet, Pune\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11.5)
    r.bold = True

    r = p.add_run("An Autonomous Institute | NBA Accredited (4 UG Programs) | NAAC A++ Accredited | ISO 21001:2018 Certified\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(8.5)

    r = p.add_run("IQAC PCCOER")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r.bold = True

    c02 = header_table.cell(0, 2)
    c02.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = c02.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if pcet_logo and os.path.exists(pcet_logo):
        p.add_run().add_picture(pcet_logo, width=Inches(0.85))

    # Row 1: Academic Year | UNIT TEST | Record No.
    c10 = header_table.cell(1, 0)
    c10.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = c10.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run("Academic Year:\n")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10)
    r = p.add_run(f"{paper_data.get('academic_year', '2025 – 26')}\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10)
    r = p.add_run("Term: ")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10)
    r = p.add_run(f"{paper_data.get('term', 'II')}")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10)

    c11 = header_table.cell(1, 1)
    c11.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = c11.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{paper_data.get('exam_type', 'UNIT TEST')}")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(16)

    c12 = header_table.cell(1, 2)
    c12.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = c12.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run("Record No.:\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10.5)
    r = p.add_run(f"{paper_data.get('record_no', 'ACAD/R/11')}")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    for r_idx in range(2):
        for c_idx in range(3):
            cell = header_table.cell(r_idx, c_idx)
            set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="000000", sz="6")
            set_cell_margins(cell, top=14, bottom=14, left=25, right=25)

    # 1.1 Space between Header and Information (department, class, div)
    p_sp1 = doc.add_paragraph()
    p_sp1.paragraph_format.space_before = Pt(7)
    p_sp1.paragraph_format.space_after = Pt(0)
    p_sp1.paragraph_format.line_spacing = Pt(1)

    # 2. Metadata Section (Department, Subject, Marks, etc.)
    meta_table = doc.add_table(rows=3, cols=3)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_widths = [Inches(3.3), Inches(2.2), Inches(1.97)]
    for row in meta_table.rows:
        for idx, w in enumerate(meta_widths):
            row.cells[idx].width = w

    # Row 0
    c = meta_table.cell(0, 0)
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run("Department: ")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11.5)
    r = p.add_run(paper_data.get('department', 'Computer Engineering'))
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11.5)

    c = meta_table.cell(0, 1)
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run("Class: ")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r = p.add_run(paper_data.get('class', 'TE'))
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    c = meta_table.cell(0, 2)
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run("Div.: ")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r = p.add_run(paper_data.get('div', 'A, B, C, D, E, F'))
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    # Row 1
    c = meta_table.cell(1, 0)
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run("Subject: ")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r = p.add_run(paper_data.get('subject', ''))
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    c = meta_table.cell(1, 1)
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run("Maximum Marks: ")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r = p.add_run(str(paper_data.get('max_marks', '30')))
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    c = meta_table.cell(1, 2)
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run("Duration: ")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r = p.add_run(str(paper_data.get('duration', '60 Min')))
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    # Row 2
    c = meta_table.cell(2, 0)
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run("Subject Code: ")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r = p.add_run(paper_data.get('subject_code', ''))
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    c = meta_table.cell(2, 1)
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)

    c = meta_table.cell(2, 2)
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run("Date: ")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r = p.add_run(paper_data.get('date', ''))
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    for row in meta_table.rows:
        for cell in row.cells:
            set_cell_borders(cell, top=None, bottom=None, left=None, right=None)
            set_cell_margins(cell, top=6, bottom=6, left=20, right=20)

    # 3. Notes
    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_before = Pt(10)
    p_note.paragraph_format.space_after = Pt(2)
    p_note.paragraph_format.line_spacing = 1.0
    r = p_note.add_run("Note:  1. Solve Que.1 or Que.2 and Que.3 or Que.4.\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10.5)
    r.italic = True
    r.bold = True
    r = p_note.add_run("          2. Give explanation or justification wherever required.")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10.5)
    r.italic = True

    # 4. Course Outcomes Table
    p_co_label = doc.add_paragraph()
    p_co_label.paragraph_format.space_before = Pt(10)
    p_co_label.paragraph_format.space_after = Pt(4)
    p_co_label.paragraph_format.line_spacing = 1.0
    r = p_co_label.add_run("Course Outcomes:")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    co_list = paper_data.get('co_list', [])
    if not co_list:
        co_list = [
            {'co': 'CO1', 'desc': 'Understand the fundamental concepts and principles of the subject.', 'bt': 'L2'},
            {'co': 'CO2', 'desc': 'Apply analytical techniques and methodologies to solve domain problems.', 'bt': 'L3'},
            {'co': 'CO3', 'desc': 'Design and evaluate architectural models and practical solutions.', 'bt': 'L3'},
        ]

    co_table = doc.add_table(rows=1 + len(co_list), cols=3)
    co_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    co_widths = [Inches(0.85), Inches(5.42), Inches(1.2)]
    for row in co_table.rows:
        for idx, w in enumerate(co_widths):
            row.cells[idx].width = w

    headers = ["CO", "Course Outcomes", "BT Level"]
    for idx, text in enumerate(headers):
        cell = co_table.cell(0, idx)
        set_cell_shading(cell, "D9D9D9")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text)
        r.bold = True
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="000000", sz="6")
        set_cell_margins(cell, top=12, bottom=12, left=25, right=25)

    for r_idx, co_item in enumerate(co_list):
        cell_co = co_table.cell(r_idx + 1, 0)
        p = cell_co.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(co_item.get('co', f'CO{r_idx+1}'))
        r.bold = True
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10.5)
        set_cell_borders(cell_co, top="single", bottom="single", left="single", right="single", color="000000", sz="6")
        set_cell_margins(cell_co, top=12, bottom=12, left=25, right=25)

        cell_desc = co_table.cell(r_idx + 1, 1)
        p = cell_desc.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(co_item.get('desc', ''))
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10.5)
        set_cell_borders(cell_desc, top="single", bottom="single", left="single", right="single", color="000000", sz="6")
        set_cell_margins(cell_desc, top=12, bottom=12, left=25, right=25)

        cell_bt = co_table.cell(r_idx + 1, 2)
        p = cell_bt.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(co_item.get('bt', 'L2'))
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10.5)
        set_cell_borders(cell_bt, top="single", bottom="single", left="single", right="single", color="000000", sz="6")
        set_cell_margins(cell_bt, top=12, bottom=12, left=25, right=25)

    # 5. Main Questions Table (Directly added, compact spacing!)
    raw_questions = paper_data.get('questions', [])
    normalized_q = []
    if raw_questions and isinstance(raw_questions[0], dict) and 'que' in raw_questions[0]:
        normalized_q = raw_questions
    elif paper_data.get('sections'):
        import re
        sections = paper_data.get('sections', [])
        for s_idx, sec in enumerate(sections):
            is_or = s_idx > 0 and (
                bool(sec.get('is_or_choice') or sec.get('isOrChoice')) or
                bool(re.search(r'\bOR\b', sec.get('name', ''), re.IGNORECASE)) or
                (bool(paper_data.get('use_or_pairing') or paper_data.get('has_or_choices')) and s_idx % 2 == 1)
            )
            if is_or:
                normalized_q.append({'type': 'OR'})

            m = re.search(r'Que\s*(\d+)', sec.get('name', ''), re.IGNORECASE)
            que_num_str = str(m.group(1)) if m else str(s_idx + 1)

            sec_qs = sec.get('questions', [])
            for q_idx, q in enumerate(sec_qs):
                sub_label = chr(65 + q_idx)
                btl = q.get('blooms_level', 'L2').split('_')[0] if q.get('blooms_level') else 'L2'
                co_val = q.get('co') or f"CO{(s_idx % 3) + 1}"
                marks_val = q.get('marks', 5)
                normalized_q.append({
                    'que': que_num_str,
                    'sub': sub_label,
                    'text': q.get('text', ''),
                    'marks_co_btl': f"{marks_val} / {co_val} / {btl}",
                    'pi': q.get('pi', '1.1.1')
                })
    if not normalized_q:
        # Fallback default questions matching standard PCCOER format
        normalized_q = [
            {'que': '1', 'sub': 'A', 'text': 'Explain the three-tier database architecture with a neat schematic diagram.', 'marks_co_btl': '5 / CO1 / L2', 'pi': '1.3.1'},
            {'que': '1', 'sub': 'B', 'text': 'Differentiate between File Processing System and Database Management System.', 'marks_co_btl': '5 / CO1 / L2', 'pi': '1.3.1'},
            {'que': '1', 'sub': 'C', 'text': 'State the roles and responsibilities of a Database Administrator (DBA).', 'marks_co_btl': '5 / CO1 / L1', 'pi': '1.2.1'},
            {'type': 'OR'},
            {'que': '2', 'sub': 'A', 'text': 'Define Data Independence. Explain Logical and Physical Data Independence with examples.', 'marks_co_btl': '5 / CO1 / L2', 'pi': '1.3.1'},
            {'que': '2', 'sub': 'B', 'text': 'Construct an ER diagram for a Hospital Management System with appropriate entities, attributes, and relationships.', 'marks_co_btl': '5 / CO2 / L3', 'pi': '2.1.2'},
            {'que': '2', 'sub': 'C', 'text': 'Explain generalization, specialization, and aggregation with suitable illustrations.', 'marks_co_btl': '5 / CO2 / L2', 'pi': '2.1.1'},
            {'que': '3', 'sub': 'A', 'text': 'What is Relational Algebra? Explain fundamental relational algebra operations with examples.', 'marks_co_btl': '5 / CO3 / L2', 'pi': '1.4.1'},
            {'que': '3', 'sub': 'B', 'text': 'Given relation schema Employee(eid, ename, salary, did), write relational algebra expressions to retrieve employees with salary > 50000.', 'marks_co_btl': '5 / CO3 / L3', 'pi': '2.2.1'},
            {'que': '3', 'sub': 'C', 'text': 'Explain Primary Key, Foreign Key, Candidate Key, and Super Key with schema examples.', 'marks_co_btl': '5 / CO2 / L2', 'pi': '1.2.2'},
            {'type': 'OR'},
            {'que': '4', 'sub': 'A', 'text': 'Explain various DDL and DML commands in SQL with their syntax and sample queries.', 'marks_co_btl': '5 / CO3 / L2', 'pi': '1.4.1'},
            {'que': '4', 'sub': 'B', 'text': 'Discuss aggregate functions in SQL (COUNT, SUM, AVG, MIN, MAX) along with GROUP BY and HAVING clauses.', 'marks_co_btl': '5 / CO3 / L3', 'pi': '2.2.2'},
            {'que': '4', 'sub': 'C', 'text': 'Illustrate nested subqueries and correlated subqueries with suitable relational queries.', 'marks_co_btl': '5 / CO3 / L3', 'pi': '2.2.2'},
        ]
    # Spacing between Course Outcomes and Questions Table
    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_before = Pt(10)
    p_sp2.paragraph_format.space_after = Pt(0)
    p_sp2.paragraph_format.line_spacing = Pt(1)

    q_table = doc.add_table(rows=1, cols=5)
    q_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    q_widths = [Inches(0.48), Inches(0.55), Inches(4.74), Inches(1.05), Inches(0.65)]
    for idx, w in enumerate(q_widths):
        q_table.rows[0].cells[idx].width = w

    q_headers = ["Que", "Sub\nQue.", "Questions", "Marks/\nCO/BTL", "PI"]
    for idx, text in enumerate(q_headers):
        cell = q_table.cell(0, idx)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text)
        r.bold = True
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="000000", sz="6")
        set_cell_margins(cell, top=18, bottom=18, left=25, right=25)

    # Track row ranges for each question group to merge column 0 vertically
    que_row_groups = []
    current_que_num = None
    group_start_idx = None

    for item in normalized_q:
        row = q_table.add_row()
        make_row_cant_split(row)
        current_row_idx = len(q_table.rows) - 1
        for idx, w in enumerate(q_widths):
            row.cells[idx].width = w

        if item.get('type') == 'OR':
            if current_que_num is not None and group_start_idx is not None:
                que_row_groups.append((current_que_num, group_start_idx, current_row_idx - 1))
                current_que_num = None
                group_start_idx = None

            merged = row.cells[0].merge(row.cells[1]).merge(row.cells[2]).merge(row.cells[3]).merge(row.cells[4])
            merged.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = merged.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run("OR")
            r.bold = True
            r.font.name = 'Times New Roman'
            r.font.size = Pt(11)
            set_cell_borders(merged, top="single", bottom="single", left="single", right="single", color="000000", sz="6")
            set_cell_margins(merged, top=12, bottom=12, left=25, right=25)
        else:
            item_que = item.get('que', '')
            if item_que != current_que_num:
                if current_que_num is not None and group_start_idx is not None:
                    que_row_groups.append((current_que_num, group_start_idx, current_row_idx - 1))
                current_que_num = item_que
                group_start_idx = current_row_idx

            c0 = row.cells[0]
            c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            c1 = row.cells[1]
            c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p1 = c1.paragraphs[0]
            p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p1.paragraph_format.space_before = Pt(0)
            p1.paragraph_format.space_after = Pt(0)
            p1.paragraph_format.line_spacing = 1.0
            r1 = p1.add_run(item.get('sub', ''))
            r1.bold = True
            r1.font.name = 'Times New Roman'
            r1.font.size = Pt(12)

            c2 = row.cells[2]
            p2 = c2.paragraphs[0]
            p2.paragraph_format.space_before = Pt(0)
            p2.paragraph_format.space_after = Pt(0)
            p2.paragraph_format.line_spacing = 1.15
            r2 = p2.add_run(item.get('text', ''))
            r2.font.name = 'Times New Roman'
            r2.font.size = Pt(12)

            c3 = row.cells[3]
            c3.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p3 = c3.paragraphs[0]
            p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p3.paragraph_format.space_before = Pt(0)
            p3.paragraph_format.space_after = Pt(0)
            p3.paragraph_format.line_spacing = 1.0
            
            # Format 3 lines: (Marks-05), [P251.1], BT6
            marks_items = parse_marks_co_btl(item.get('marks_co_btl', ''))
            for m_idx, (m_text, is_bold) in enumerate(marks_items):
                if m_idx > 0:
                    r = p3.add_run('\n')
                r = p3.add_run(m_text)
                r.bold = is_bold
                r.font.name = 'Times New Roman'
                r.font.size = Pt(11)

            c4 = row.cells[4]
            c4.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p4 = c4.paragraphs[0]
            p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p4.paragraph_format.space_before = Pt(0)
            p4.paragraph_format.space_after = Pt(0)
            p4.paragraph_format.line_spacing = 1.0
            r4 = p4.add_run(str(item.get('pi', '1.1.1')))
            r4.font.name = 'Times New Roman'
            r4.font.size = Pt(10.5)

            for cell in row.cells:
                set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="000000", sz="6")
                set_cell_margins(cell, top=20, bottom=20, left=25, right=25)

    # Finalize last question group if any
    if current_que_num is not None and group_start_idx is not None:
        que_row_groups.append((current_que_num, group_start_idx, len(q_table.rows) - 1))

    # Vertically merge Que column (cell 0) across all sub-questions for each question group
    # and keep all sub-questions of that question together on the same page
    for que_num, s_idx, e_idx in que_row_groups:
        start_cell = q_table.cell(s_idx, 0)
        if s_idx < e_idx:
            end_cell = q_table.cell(e_idx, 0)
            merged_que_cell = start_cell.merge(end_cell)
            for r_i in range(s_idx, e_idx):
                for cell in q_table.rows[r_i].cells:
                    for cp in cell.paragraphs:
                        cp.paragraph_format.keep_with_next = True
        else:
            merged_que_cell = start_cell

        merged_que_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = merged_que_cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.text = ""
        r = p.add_run(str(que_num))
        r.bold = True
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        set_cell_borders(merged_que_cell, top="single", bottom="single", left="single", right="single", color="000000", sz="6")
        set_cell_margins(merged_que_cell, top=20, bottom=20, left=25, right=25)

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream
